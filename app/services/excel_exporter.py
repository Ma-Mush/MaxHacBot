"""Professional Excel workbook generator with corporate styling and openpyxl."""
import io
import logging
from typing import Any, Dict, List, Optional, Sequence, Union
import openpyxl
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Corporate Style Constants
COLOR_HEADER_BG = "1E293B"      # Deep Slate Navy
COLOR_HEADER_TEXT = "FFFFFF"    # White
COLOR_SUBHEADER_BG = "334155"   # Slate 700
COLOR_ZEBRA = "F8FAFC"          # Slate 50
COLOR_CARD_BG = "F1F5F9"        # Slate 100
COLOR_CARD_BORDER = "CBD5E1"    # Slate 300
COLOR_BORDER = "E2E8F0"         # Slate 200
COLOR_POSITIVE = "10B981"       # Emerald Green
COLOR_NEGATIVE = "EF4444"       # Rose Red
COLOR_MUTED = "64748B"          # Slate 500

FONT_NAME = "Segoe UI"


class ExcelExporterService:
    """Service to create branded, presentation-ready Excel workbooks."""

    def __init__(self) -> None:
        self.font_title = Font(name=FONT_NAME, size=16, bold=True, color="0F172A")
        self.font_subtitle = Font(name=FONT_NAME, size=10, italic=True, color=COLOR_MUTED)
        self.font_header = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_HEADER_TEXT)
        self.font_bold = Font(name=FONT_NAME, size=10, bold=True)
        self.font_body = Font(name=FONT_NAME, size=10)

        self.fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        self.fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
        self.fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        self.fill_card = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")

        self.thin_border = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
        self.card_border = Border(
            left=Side(style="medium", color=COLOR_CARD_BORDER),
            right=Side(style="medium", color=COLOR_CARD_BORDER),
            top=Side(style="medium", color=COLOR_CARD_BORDER),
            bottom=Side(style="medium", color=COLOR_CARD_BORDER),
        )

    def create_workbook(self) -> openpyxl.Workbook:
        """Create a new workbook with default gridlines enabled."""
        wb = openpyxl.Workbook()
        # Enable grid lines on initial sheet
        ws = wb.active
        ws.views.sheetView[0].showGridLines = True
        return wb

    def add_title_block(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        title: str,
        subtitle: Optional[str] = None,
        start_row: int = 1,
    ) -> int:
        """Add branded report title and subtitle block."""
        ws.cell(row=start_row, column=1, value=title).font = self.font_title
        current_row = start_row + 1
        if subtitle:
            ws.cell(row=current_row, column=1, value=subtitle).font = self.font_subtitle
            current_row += 1
        return current_row + 1  # Empty row after title

    def add_kpi_cards(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        kpis: List[Dict[str, Any]],
        start_row: int = 4,
        start_col: int = 1,
    ) -> int:
        """Add styled executive KPI cards side-by-side."""
        col_offset = start_col
        for kpi in kpis:
            title = kpi.get("title", "").upper()
            value = kpi.get("value", "")
            delta = kpi.get("delta")
            is_pos = kpi.get("is_positive")

            # Title cell
            cell_t = ws.cell(row=start_row, column=col_offset, value=title)
            cell_t.font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_MUTED)
            cell_t.fill = self.fill_card
            cell_t.alignment = Alignment(horizontal="center", vertical="center")

            # Value cell
            cell_v = ws.cell(row=start_row + 1, column=col_offset, value=value)
            cell_v.font = Font(name=FONT_NAME, size=15, bold=True, color="0F172A")
            cell_v.fill = self.fill_card
            cell_v.alignment = Alignment(horizontal="center", vertical="center")

            # Delta cell
            delta_color = COLOR_POSITIVE if is_pos else (COLOR_NEGATIVE if is_pos is False else COLOR_MUTED)
            prefix = "▲ " if is_pos else ("▼ " if is_pos is False else "")
            delta_val = f"{prefix}{delta}" if delta else "—"
            cell_d = ws.cell(row=start_row + 2, column=col_offset, value=delta_val)
            cell_d.font = Font(name=FONT_NAME, size=9, bold=True, color=delta_color)
            cell_d.fill = self.fill_card
            cell_d.alignment = Alignment(horizontal="center", vertical="center")

            # Apply borders to card cells
            for r in range(start_row, start_row + 3):
                ws.cell(row=r, column=col_offset).border = self.card_border

            # Width for KPI column
            ws.column_dimensions[get_column_letter(col_offset)].width = 22
            col_offset += 2  # Space out cards

        return start_row + 4

    def add_table(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        headers: List[str],
        rows: List[List[Any]],
        start_row: int = 8,
        start_col: int = 1,
        number_formats: Optional[Dict[int, str]] = None,
    ) -> int:
        """Add styled corporate table with header, zebra striping, and number formatting."""
        # Header row
        for col_idx, header in enumerate(headers, start=start_col):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = self.thin_border
        ws.row_dimensions[start_row].height = 24

        # Data rows
        current_row = start_row + 1
        for row_data in rows:
            is_even = (current_row % 2 == 0)
            row_fill = self.fill_zebra if is_even else self.fill_white

            for col_idx, val in enumerate(row_data, start=start_col):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = self.font_body
                cell.fill = row_fill
                cell.border = self.thin_border

                # Alignment and Number format
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                if number_formats and col_idx in number_formats:
                    cell.number_format = number_formats[col_idx]

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        self.auto_fit_columns(ws, min_col=start_col, max_col=start_col + len(headers) - 1)
        return current_row + 1

    def embed_chart_image(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        image_bytes: bytes,
        cell_ref: str = "A1",
        width: int = 680,
        height: int = 340,
    ) -> None:
        """Embed chart PNG image into specific worksheet cell reference."""
        try:
            img_io = io.BytesIO(image_bytes)
            img = OpenPyxlImage(img_io)
            img.width = width
            img.height = height
            ws.add_image(img, cell_ref)
        except Exception as exc:
            logger.error(f"Failed to embed chart image into Excel at {cell_ref}: {exc}")

    def auto_fit_columns(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        min_col: int = 1,
        max_col: Optional[int] = None,
        padding: int = 4,
    ) -> None:
        """Dynamically adjust column widths to avoid text truncation."""
        if max_col is None:
            max_col = ws.max_column

        for col in range(min_col, max_col + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            calculated_width = max(max_len + padding, 14)
            # Cap width so comments/long text do not make huge columns
            ws.column_dimensions[col_letter].width = min(calculated_width, 45)

    def to_bytes(self, wb: openpyxl.Workbook) -> bytes:
        """Serialize workbook to byte buffer."""
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()


excel_exporter = ExcelExporterService()
